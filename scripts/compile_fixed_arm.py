from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from inspect_ai.log import read_eval_log

ROOT = Path("experiment_runs")
OUTPUT_DIR = ROOT / "compiled"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
timestamped_workbook = (
    OUTPUT_DIR / f"fixed_arm_master_{timestamp}.xlsx"
)
latest_workbook = OUTPUT_DIR / "fixed_arm_master_results.xlsx"
runs_csv_output = OUTPUT_DIR / "fixed_arm_master_runs.csv"

required_levels = set(range(1, 13))
selected_batches: dict[int, Path] = {}
audit_rows = []


def get_level(manifest: dict, batch_path: Path):
    value = manifest.get("specificity_level")

    if value is not None:
        return int(value)

    marker = "_L"

    if marker in batch_path.name:
        return int(batch_path.name.split(marker, 1)[1][:2])

    return None


def inspect_batch(batch_path: Path):
    manifest_path = batch_path / "batch_manifest.json"

    audit = {
        "batch_id": batch_path.name,
        "batch_path": str(batch_path.resolve()),
        "level": None,
        "selected": False,
        "eligible": False,
        "recorded_runs": 0,
        "completed_runs": 0,
        "reason": "",
    }

    if not manifest_path.is_file():
        audit["reason"] = "Missing batch_manifest.json"
        return audit

    try:
        manifest = json.loads(
            manifest_path.read_text(encoding="utf-8")
        )
    except Exception as error:
        audit["reason"] = f"Manifest read error: {error}"
        return audit

    level = get_level(manifest, batch_path)
    runs = manifest.get("runs", [])

    completed_runs = [
        run
        for run in runs
        if run.get("status") == "completed"
        and run.get("report_exit_code", 0) == 0
    ]

    audit["level"] = level
    audit["recorded_runs"] = len(runs)
    audit["completed_runs"] = len(completed_runs)

    reasons = []

    if manifest.get("arm") != "fixed":
        reasons.append("not fixed arm")

    if level not in required_levels:
        reasons.append("invalid specificity level")

    if manifest.get("planned_repeats") != 3:
        reasons.append("planned repeats is not 3")

    if manifest.get("turn_budget") != 150:
        reasons.append("turn budget is not 150")

    if len(completed_runs) != 3:
        reasons.append(
            f"{len(completed_runs)} completed runs instead of 3"
        )

    runs_csv = batch_path / "batch_report" / "runs.csv"

    if not runs_csv.is_file():
        reasons.append("missing batch_report/runs.csv")
    else:
        try:
            row_count = len(pd.read_csv(runs_csv))

            if row_count != 3:
                reasons.append(
                    f"runs.csv contains {row_count} rows instead of 3"
                )
        except Exception as error:
            reasons.append(f"runs.csv read error: {error}")

    if reasons:
        audit["reason"] = "; ".join(reasons)
        return audit

    audit["eligible"] = True
    audit["reason"] = "eligible formal batch"
    return audit


for batch_path in sorted(ROOT.glob("FLE_fixed_L??_B*")):
    if not batch_path.is_dir():
        continue

    audit = inspect_batch(batch_path)
    audit_rows.append(audit)

    if audit["eligible"]:
        level = audit["level"]

        if (
            level not in selected_batches
            or batch_path.name > selected_batches[level].name
        ):
            selected_batches[level] = batch_path


missing_levels = sorted(
    required_levels - set(selected_batches)
)

if missing_levels:
    print("COMPILATION STOPPED")
    print("=" * 70)
    print(
        "No eligible three-run formal batch was found for levels:",
        missing_levels,
    )
    print()
    print("Batch audit:")

    for audit in audit_rows:
        print(
            f"{audit['batch_id']}: "
            f"level={audit['level']}; "
            f"eligible={audit['eligible']}; "
            f"{audit['reason']}"
        )

    raise SystemExit(1)


for audit in audit_rows:
    level = audit["level"]

    if (
        audit["eligible"]
        and level in selected_batches
        and audit["batch_id"] == selected_batches[level].name
    ):
        audit["selected"] = True
        audit["reason"] = "selected formal batch"
    elif audit["eligible"]:
        audit["reason"] = (
            "eligible but superseded by a newer formal batch"
        )


def read_batch_table(
    batch_path: Path,
    filename: str,
    level: int,
):
    path = batch_path / "batch_report" / filename

    if not path.is_file():
        return pd.DataFrame()

    frame = pd.read_csv(path)
    frame.insert(0, "source_batch", batch_path.name)

    normalized_columns = {
        str(column).lower().replace(" ", "_").replace("-", "_")
        for column in frame.columns
    }

    if (
        "specificity_level" not in normalized_columns
        and "level" not in normalized_columns
    ):
        frame.insert(1, "specificity_level", level)

    return frame


table_definitions = {
    "Runs": "runs.csv",
    "Behaviors": "behaviors.csv",
    "Errors": "errors.csv",
    "Phases": "phases.csv",
    "Turns": "turns.csv",
}

combined_tables = {}

for sheet_name, filename in table_definitions.items():
    frames = []

    for level in sorted(selected_batches):
        frame = read_batch_table(
            selected_batches[level],
            filename,
            level,
        )

        if not frame.empty:
            frames.append(frame)

    combined_tables[sheet_name] = (
        pd.concat(frames, ignore_index=True, sort=False)
        if frames
        else pd.DataFrame()
    )


runs = combined_tables["Runs"].copy()

if len(runs) != 36:
    raise SystemExit(
        f"Expected 36 formal runs but compiled {len(runs)}."
    )


def normalize_name(value):
    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
    )


def find_column(frame, candidates):
    lookup = {
        normalize_name(column): column
        for column in frame.columns
    }

    for candidate in candidates:
        if candidate in lookup:
            return lookup[candidate]

    return None


level_column = find_column(
    runs,
    ["specificity_level", "level"],
)

production_column = find_column(
    runs,
    [
        "production_score",
        "final_score",
        "production",
    ],
)

automation_column = find_column(
    runs,
    [
        "automated_production_score",
        "final_automated_score",
        "automation_score",
    ],
)

technology_column = find_column(
    runs,
    [
        "technologies",
        "technology_count",
        "researched_technologies",
    ],
)

required_columns = {
    "specificity level": level_column,
    "production": production_column,
    "automation": automation_column,
    "technologies": technology_column,
}

missing_columns = [
    name
    for name, column in required_columns.items()
    if column is None
]

if missing_columns:
    print("Available runs.csv columns:")
    print(list(runs.columns))

    raise SystemExit(
        "Missing required columns: "
        + ", ".join(missing_columns)
    )


runs["fdi_production_raw"] = pd.to_numeric(
    runs[production_column],
    errors="coerce",
)

runs["fdi_automation_raw"] = pd.to_numeric(
    runs[automation_column],
    errors="coerce",
)

def count_technologies(value):
    if pd.isna(value):
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    names = [
        item.strip()
        for item in str(value).split(",")
        if item.strip()
    ]

    return float(len(names))


def numeric_score_value(score):
    value = getattr(score, "value", score)

    if isinstance(value, dict):
        for key in ("value", "score", "mean"):
            if key in value:
                value = value[key]
                break

    if isinstance(value, (list, tuple)) and len(value) == 1:
        value = value[0]

    return float(value)


def read_achievement_from_eval(run_directory):
    eval_files = sorted(
        path
        for path in run_directory.glob("*.eval")
        if "/raw/" not in str(path).replace("\\", "/")
    )

    if not eval_files:
        raw_directory = run_directory / "raw"

        if raw_directory.is_dir():
            eval_files = sorted(
                raw_directory.glob("*.eval")
            )

    if not eval_files:
        raise RuntimeError(
            f"No evaluation file found in {run_directory}"
        )

    log = read_eval_log(str(eval_files[-1]))

    if not log.samples:
        raise RuntimeError(
            f"No samples found in {eval_files[-1]}"
        )

    scores = log.samples[0].scores or {}

    for score_name, score in scores.items():
        if "achievement" in str(score_name).lower():
            return numeric_score_value(score)

    raise RuntimeError(
        "Achievement scorer was not found in "
        f"{eval_files[-1]}. Available scorers: "
        f"{list(scores.keys())}"
    )


def read_achievement_from_console(console_path):
    if not console_path.is_file():
        raise RuntimeError(
            f"Missing console log: {console_path}"
        )

    text = console_path.read_text(
        encoding="utf-8",
        errors="replace",
    )

    text = re.sub(
        r"\x1b\[[0-9;]*[A-Za-z]",
        "",
        text,
    )

    pattern = re.compile(
        r"production_score\s+"
        r"technologies\s+"
        r"achievements\s+"
        r"automated_production_score"
        r".*?"
        r"mean\s+([0-9,]+(?:\.[0-9]+)?)\s+"
        r"mean\s+([0-9,]+(?:\.[0-9]+)?)\s+"
        r"mean\s+([0-9,]+(?:\.[0-9]+)?)\s+"
        r"mean\s+([0-9,]+(?:\.[0-9]+)?)",
        re.DOTALL,
    )

    matches = pattern.findall(text)

    if not matches:
        raise RuntimeError(
            f"No scorer table found in {console_path}"
        )

    return float(
        matches[-1][2].replace(",", "")
    )


def read_achievement_score(row):
    run_directory = (
        ROOT
        / str(row["source_batch"])
        / str(row["run_id"])
    )

    eval_error = None

    try:
        return read_achievement_from_eval(
            run_directory
        )
    except Exception as error:
        eval_error = error

    console_path = run_directory / "console.log"

    try:
        return read_achievement_from_console(
            console_path
        )
    except Exception as console_error:
        raise RuntimeError(
            "Could not recover the achievement score for "
            f"{row['run_id']}.\n"
            f"Eval error: {eval_error}\n"
            f"Console error: {console_error}"
        ) from console_error


runs["fdi_technologies_raw"] = (
    runs[technology_column].apply(count_technologies)
)

runs["fdi_achievements_raw"] = runs.apply(
    read_achievement_score,
    axis=1,
)

score_columns = [
    "fdi_production_raw",
    "fdi_automation_raw",
    "fdi_technologies_raw",
    "fdi_achievements_raw",
]

if runs[score_columns].isna().any().any():
    problem_rows = runs[
        runs[score_columns].isna().any(axis=1)
    ]

    print("Rows containing missing or nonnumeric scores:")
    print(problem_rows.to_string(index=False))

    raise SystemExit(
        "FDI calculation stopped because one or more scores "
        "were missing or nonnumeric."
    )


def zscore(series):
    series = pd.to_numeric(series, errors="coerce")
    standard_deviation = series.std(ddof=1)

    if (
        pd.isna(standard_deviation)
        or standard_deviation == 0
    ):
        return pd.Series(0.0, index=series.index)

    return (
        series - series.mean()
    ) / standard_deviation


runs["fdi_production_z"] = zscore(
    np.log1p(runs["fdi_production_raw"])
)

runs["fdi_automation_z"] = zscore(
    np.log1p(runs["fdi_automation_raw"])
)

runs["fdi_technology_z"] = zscore(
    runs["fdi_technologies_raw"]
)

runs["fdi_achievement_z"] = zscore(
    runs["fdi_achievements_raw"]
)

runs["fdi_advancement_z"] = (
    runs["fdi_technology_z"]
    + runs["fdi_achievement_z"]
) / 2

runs["fdi_z"] = (
    runs["fdi_production_z"]
    + runs["fdi_automation_z"]
    + runs["fdi_advancement_z"]
) / 3

runs["factory_development_index"] = (
    50 + 10 * runs["fdi_z"]
)

runs[level_column] = pd.to_numeric(
    runs[level_column],
    errors="raise",
).astype(int)

runs = runs.sort_values(
    [level_column, "source_batch"]
).reset_index(drop=True)

level_summary = (
    runs.groupby(level_column)
    .agg(
        runs=("factory_development_index", "size"),
        production_mean=("fdi_production_raw", "mean"),
        production_sd=("fdi_production_raw", "std"),
        automation_mean=("fdi_automation_raw", "mean"),
        automation_sd=("fdi_automation_raw", "std"),
        technologies_mean=("fdi_technologies_raw", "mean"),
        achievements_mean=("fdi_achievements_raw", "mean"),
        fdi_mean=("factory_development_index", "mean"),
        fdi_sd=("factory_development_index", "std"),
        fdi_min=("factory_development_index", "min"),
        fdi_max=("factory_development_index", "max"),
    )
    .reset_index()
    .sort_values(level_column)
)

batch_audit = pd.DataFrame(audit_rows).sort_values(
    ["level", "batch_id"],
    na_position="last",
)

fdi_method = pd.DataFrame(
    [
        {
            "component": "Production",
            "transformation": (
                "log(1 + production score), then pooled z-score"
            ),
            "weight": "33.33%",
        },
        {
            "component": "Automation",
            "transformation": (
                "log(1 + automated production score), "
                "then pooled z-score"
            ),
            "weight": "33.33%",
        },
        {
            "component": "Advancement",
            "transformation": (
                "mean of technology and achievement "
                "pooled z-scores"
            ),
            "weight": "33.33%",
        },
        {
            "component": "Reported FDI",
            "transformation": (
                "50 + 10 × mean of the three component z-scores"
            ),
            "weight": "Composite",
        },
    ]
)

combined_tables["Runs"] = runs

with pd.ExcelWriter(
    timestamped_workbook,
    engine="openpyxl",
) as writer:
    level_summary.to_excel(
        writer,
        sheet_name="Level Summary",
        index=False,
    )

    runs.to_excel(
        writer,
        sheet_name="Runs",
        index=False,
    )

    for sheet_name in (
        "Behaviors",
        "Errors",
        "Phases",
        "Turns",
    ):
        frame = combined_tables[sheet_name]

        if not frame.empty:
            frame.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    batch_audit.to_excel(
        writer,
        sheet_name="Batch Audit",
        index=False,
    )

    fdi_method.to_excel(
        writer,
        sheet_name="FDI Method",
        index=False,
    )

    workbook = writer.book

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        for column_cells in worksheet.columns:
            values = [
                str(cell.value)
                for cell in column_cells
                if cell.value is not None
            ]

            width = min(
                max(
                    [len(value) for value in values] + [10]
                ) + 2,
                45,
            )

            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = width

    summary_sheet = workbook["Level Summary"]

    headers = {
        cell.value: cell.column
        for cell in summary_sheet[1]
    }

    chart = LineChart()
    chart.title = "Fixed-Arm Performance by Specificity Level"
    chart.style = 2
    chart.y_axis.title = "Mean score"
    chart.x_axis.title = "Specificity level"
    chart.height = 9
    chart.width = 17

    production_reference = Reference(
        summary_sheet,
        min_col=headers["production_mean"],
        min_row=1,
        max_row=13,
    )

    automation_reference = Reference(
        summary_sheet,
        min_col=headers["automation_mean"],
        min_row=1,
        max_row=13,
    )

    category_reference = Reference(
        summary_sheet,
        min_col=headers[level_column],
        min_row=2,
        max_row=13,
    )

    chart.add_data(
        production_reference,
        titles_from_data=True,
    )

    chart.add_data(
        automation_reference,
        titles_from_data=True,
    )

    chart.set_categories(category_reference)
    summary_sheet.add_chart(chart, "M2")

runs.to_csv(
    runs_csv_output,
    index=False,
)

shutil.copy2(
    timestamped_workbook,
    latest_workbook,
)

print()
print("FIXED-ARM COMPILATION COMPLETE")
print("=" * 72)
print(f"Selected formal batches: {len(selected_batches)}")
print(f"Formal runs compiled: {len(runs)}")
print()

print("Selected batches:")

for level in sorted(selected_batches):
    print(
        f"  Level {level:02d}: "
        f"{selected_batches[level].name}"
    )

print()
print("Master workbook:")
print(timestamped_workbook.resolve())

print()
print("Stable latest workbook:")
print(latest_workbook.resolve())

print()
print("Master run-level CSV:")
print(runs_csv_output.resolve())

print()
print("Level summary:")
print(
    level_summary[
        [
            level_column,
            "production_mean",
            "automation_mean",
            "technologies_mean",
            "achievements_mean",
            "fdi_mean",
        ]
    ].round(3).to_string(index=False)
)
